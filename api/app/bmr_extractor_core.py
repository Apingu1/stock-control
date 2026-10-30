"""Offline Word BMR extraction used by the StkFlo staged migration workflow.

Legacy .doc files are converted to .docx with headless LibreOffice. Source files
are never executed; only document text and table XML are read. No stock writes.
"""
from __future__ import annotations
import csv
import hashlib
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, BadZipFile

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PRODUCT_HEADERS = [
 "product_code","product_name_bmr","product_name_product_list","bmr_version",
 "bmr_document_number","bmr_shelf_life","bmr_shelf_life_days","product_list_shelf_life_days",
 "storage_conditions","licensed_status","medisource_line","export_line","controlled_drug",
 "high_potency","stock_line","fridge_line","product_list_remarks","formula_basis",
 "material_count","review_status","bmr_file","source_product_list_row"
]
FORMULA_HEADERS = [
 "product_code","bmr_version","product_name","material_sequence","material_name",
 "material_strength","master_formula_raw","quantity_type","quantity_value","quantity_unit",
 "qs_target_value","qs_target_unit","conditional_material","manufacturer_supplier",
 "formula_basis","bmr_file","source_section","section4_table","section4_row"
]
MAX_DOCS = 1100
MAX_SOURCE_BYTES = 120_000_000


def _clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _csv(rows, keys):
    f = io.StringIO()
    writer = csv.DictWriter(f, keys, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return "\ufeff" + f.getvalue()


def _filename(name):
    match = re.search(r"([A-Z][A-Z0-9()_-]{1,30})\s*\.\s*(V\d{1,3})\b",name.upper())
    return (match.group(1),match.group(2)) if match else ("","")


def _tables(data):
    with ZipFile(io.BytesIO(data)) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    tables = []
    for table in root.findall(".//"+W+"body/"+W+"tbl"):
        tables.append([
            [_clean("".join(t.text or "" for t in cell.iter(W+"t")))
             for cell in row.findall(W+"tc")]
            for row in table.findall(W+"tr")
        ])
    return tables


def _label(tables, phrase):
    needle=re.compile(phrase,re.I)
    for table in tables:
        for row in table:
            for idx,cell in enumerate(row[:-1]):
                if needle.fullmatch(cell):
                    return row[idx+1]
    return ""


def _qty(raw):
    s = _clean(raw).lower().replace("*","")
    qs = re.match(r"^q\.?\s*s\.?\s*(?:to|up\s+to)\s*(\d+(?:\.\d+)?)\s*([a-zµμ]+)$",s)
    if qs:
        return ("QS to","", "",qs.group(1),qs.group(2))
    fixed = re.match(r"^(\d+(?:\.\d+)?)\s*([a-zµμ]+)$",s)
    if fixed:
        return ("Fixed", fixed.group(1),fixed.group(2),"","")
    return ("Unparsed" if s else "Missing","","","","")


def _product_list(payload):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
    sheet=next((s for s in wb.worksheets if s.title.casefold()=="product list"),None)
    if sheet is None:
        raise ValueError("The product list workbook has no 'Product List' sheet")
    mapping={}
    for rownum,row in enumerate(sheet.iter_rows(values_only=True),1):
        if rownum==1:
            if _clean(row[1] if len(row)>1 else "").upper()!="PRODUCT CODE":
                raise ValueError("Unexpected Product List layout (column B must contain Product Code)")
            continue
        vals=list(row)+[None]*6
        code=_clean(vals[1]).upper()
        if not code or code in mapping:
            continue
        med=_clean(vals[5]).upper()
        medisource="Yes" if med in ("M","YES") else ("No" if med in ("NO","N/A","NA") else "Unknown")
        mapping[code]={
            "product_name":_clean(vals[0]),"shelf":_clean(vals[2]),
            "licensed":_clean(vals[3]).upper(),"remarks":_clean(vals[4]),
            "medisource":medisource,"rownum":rownum,
        }
    wb.close()
    return mapping


def _parse_document(name,data,lookup,issues):
    codefile,verfile = _filename(Path(name).name)
    tables=_tables(data)
    head=tables[0][0] if tables and tables[0] else []
    title=head[0] if head else ""
    docref=""
    for cell in head:
        match=re.search(r"\b([A-Z]{1,6}\d{1,6})\s*\.\s*(V\d{1,3})\b",cell,re.I)
        if match:
            docref=match.group(1).upper()+"."+match.group(2).upper()
            break
    code=_label(tables,r"Product\s*Code").upper() or codefile
    version=docref.split(".")[-1] if docref else verfile
    shelf=_label(tables,r"Product\s+Shelf\s+Life")
    storage=_label(tables,r"Storage\s+Conditions?")
    daymatch=re.search(r"\d+",shelf)
    days=daymatch.group() if daymatch else ""
    ref=lookup.get(code,{})
    remarks=ref.get("remarks","").upper()
    flag_words=set(re.findall(r"[A-Z]+",remarks))
    rows=[]
    for ti,table in enumerate(tables,1):
        for hi,headers in enumerate(table):
            uppercase=[_clean(x).upper() for x in headers]
            if "INGREDIENT NAME" not in uppercase or "MASTER FORMULA" not in uppercase:
                continue
            nameidx=uppercase.index("INGREDIENT NAME")
            qidx=uppercase.index("MASTER FORMULA")
            strengthidx=uppercase.index("STRENGTH") if "STRENGTH" in uppercase else -1
            supplieridx=next((i for i,x in enumerate(uppercase) if "MANUFACTURER" in x and "SUPPLIER" in x),-1)
            for ri,r in enumerate(table[hi+1:],hi+2):
                ingredient=_clean(r[nameidx] if len(r)>nameidx else "")
                raw=_clean(r[qidx] if len(r)>qidx else "")
                if re.search(r"to be completed by BMR Verifier|checked by:|confirm name, strength",ingredient,re.I):
                    break
                if not ingredient or re.search(r"quantity of final ingredient|required when concentration|^\*|please ensure the api|^product\s*check",ingredient,re.I):
                    continue
                kind,quantity,unit,target,target_unit=_qty(raw)
                if kind in ("Unparsed","Missing"):
                    issues.append({"severity":"REVIEW","product_code":code,
                        "bmr_file":name,"message":f"Check {ingredient}: {raw or 'missing quantity'}"})
                rows.append({
                    "product_code":code,"bmr_version":version,"product_name":title,
                    "material_sequence":len(rows)+1,"material_name":ingredient,
                    "material_strength":r[strengthidx] if strengthidx>=0 and len(r)>strengthidx else "",
                    "master_formula_raw":raw,"quantity_type":kind,
                    "quantity_value":quantity,"quantity_unit":unit,"qs_target_value":target,
                    "qs_target_unit":target_unit,
                    "conditional_material":"Yes" if re.search(r"if\s+(required|applicable)|_{3,}",ingredient,re.I) else "No",
                    "manufacturer_supplier":r[supplieridx] if supplieridx>=0 and len(r)>supplieridx else "",
                    "formula_basis":"","bmr_file":name,"source_section":"Section 4",
                    "section4_table":ti,"section4_row":ri
                })
            break
    bases=sorted({(x["qs_target_value"],x["qs_target_unit"]) for x in rows if x["quantity_type"]=="QS to"})
    basis=", ".join(f"{value} {unit}" for value,unit in bases)
    for line in rows:
        line["formula_basis"]=basis
    product={
        "product_code":code,"product_name_bmr":title,
        "product_name_product_list":ref.get("product_name",""),
        "bmr_version":version,"bmr_document_number":docref,
        "bmr_shelf_life":shelf,"bmr_shelf_life_days":days,
        "product_list_shelf_life_days":ref.get("shelf",""),
        "storage_conditions":storage,"licensed_status":ref.get("licensed","Unknown"),
        "medisource_line":ref.get("medisource","Unknown"),
        "export_line":"EXPORT_LINE" if ref.get("medisource","").lower()=="yes" else
            ("N_A" if ref.get("medisource","").lower()=="no" else "Unknown"),
        "controlled_drug":"Yes" if "CD" in flag_words else "No",
        "high_potency":"Yes" if "HP" in flag_words else "No",
        "stock_line":"Yes" if "STOCK" in flag_words else "No",
        "fridge_line":"Yes" if "FRIDGE" in flag_words else "No",
        "product_list_remarks":ref.get("remarks",""),
        "formula_basis":basis,"material_count":len(rows),"review_status":"",
        "bmr_file":name,"source_product_list_row":ref.get("rownum","")
    }
    if not ref:
        issues.append({"severity":"REVIEW","product_code":code,
                       "bmr_file":name,"message":"Product code not found in source product list"})
    if not rows or not days or not docref:
        issues.append({"severity":"REVIEW","product_code":code,"bmr_file":name,
                       "message":"Missing Section 4 rows, shelf life or top-right BMR version"})
    if ref and days and ref.get("shelf") and days!=ref["shelf"]:
        issues.append({"severity":"REVIEW","product_code":code,"bmr_file":name,
                       "message":"Shelf life mismatch between BMR and product list"})
    product["review_status"]="REVIEW" if any(i["product_code"]==code and i["bmr_file"]==name for i in issues) else "READY FOR QA VERIFICATION"
    return product,rows


def extract(archive: bytes, product_workbook: bytes | None = None, progress=None):
    progress=progress or (lambda message: None)
    if len(archive)>MAX_SOURCE_BYTES:
        raise ValueError("BMR ZIP exceeds the 120 MB maximum")
    try:
        z=ZipFile(io.BytesIO(archive))
    except BadZipFile as exc:
        raise ValueError("Invalid BMR ZIP file") from exc
    with z:
        candidates=[]
        sheets=[]
        total=0
        for item in z.infolist():
            if item.is_dir() or Path(item.filename).name.startswith("~$"):
                continue
            ext=Path(item.filename).suffix.lower()
            if ext not in (".doc",".docx",".xlsx",".xlsm"):
                continue
            if item.file_size>12_000_000:
                raise ValueError(f"An individual file exceeds 12 MB: {item.filename}")
            total+=item.file_size
            if total>MAX_SOURCE_BYTES:
                raise ValueError("Uncompressed BMR contents exceed 120 MB")
            data=z.read(item)
            if ext in (".doc",".docx"):
                candidates.append((item.filename,data,ext))
            elif "PRODUCT LIST" in item.filename.upper():
                sheets.append(data)
    if not candidates or len(candidates)>MAX_DOCS:
        raise ValueError(f"Expected between 1 and {MAX_DOCS} Word BMR files")
    workbook=product_workbook or (sheets[0] if sheets else None)
    if not workbook:
        raise ValueError("Product List workbook is missing; upload it separately")
    lookup=_product_list(workbook)
    issues=[];products=[];formula=[]
    progress(f"Found {len(candidates)} BMRs and {len(lookup)} product list codes")
    with tempfile.TemporaryDirectory(prefix="stkflo_bmr_server_") as directory:
        root=Path(directory); originals=root/"originals"; dest=root/"converted"
        originals.mkdir();dest.mkdir()
        legacy=[]
        docx={}
        for name,data,ext in candidates:
            if ext==".docx":
                docx[name]=data
                continue
            token=hashlib.sha256(name.encode()).hexdigest()[:16]
            path=originals/(token+".doc")
            path.write_bytes(data)
            legacy.append((name,path,dest/(token+".docx")))
        if legacy:
            soffice=shutil.which("libreoffice") or shutil.which("soffice")
            if not soffice:
                raise ValueError("LibreOffice is not installed on the StkFlo API server; .doc extraction cannot run")
            for start in range(0,len(legacy),30):
                part=legacy[start:start+30]
                cmd=[soffice, "-env:UserInstallation="+(root/"lo_profile").as_uri(),
                     "--headless","--convert-to","docx","--outdir",str(dest)]+[
                     str(src) for _,src,_ in part]
                try:
                    proc=subprocess.run(cmd,timeout=240,capture_output=True,text=True)
                except subprocess.TimeoutExpired as exc:
                    raise ValueError("LibreOffice legacy BMR conversion timed out") from exc
                if proc.returncode:
                    raise ValueError("LibreOffice document conversion failed: "+proc.stderr[-200:])
                progress(f"Converted {min(start+30,len(legacy))}/{len(legacy)} legacy BMRs")
            for name,_,out in legacy:
                if out.exists():
                    docx[name]=out.read_bytes()
                else:
                    issues.append({"severity":"REVIEW","product_code":_filename(name)[0],
                                   "bmr_file":name,"message":"BMR conversion failed"})
        for index,(name,_,_) in enumerate(candidates,1):
            data=docx.get(name)
            if not data:
                continue
            try:
                p,lines=_parse_document(name,data,lookup,issues)
                if p["product_code"] and lines:
                    products.append(p)
                    formula.extend(lines)
            except Exception as exc:
                issues.append({"severity":"REVIEW","product_code":_filename(name)[0],
                               "bmr_file":name,"message":f"Extraction failed: {type(exc).__name__}"})
            if index%25==0 or index==len(candidates):
                progress(f"Extracted {index}/{len(candidates)} BMRs")
    products.sort(key=lambda x:(x["product_code"],x["bmr_version"]))
    formula.sort(key=lambda x:(x["product_code"],x["bmr_version"],x["material_sequence"]))
    return {
        "products_csv":_csv(products,PRODUCT_HEADERS),
        "formulas_csv":_csv(formula,FORMULA_HEADERS),
        "products":products,"issues":issues,
        "product_count":len(products),"formula_count":len(formula),
        "source_count":len(candidates),
    }
